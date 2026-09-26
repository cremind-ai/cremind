"""Extraction from bytes in memory: the path a Google Drive file takes.

A Drive download or export never touches the disk, so the kind has to be named
from the bytes (:func:`detect_bytes`, which reads a ZIP or OLE directory from
memory as :func:`detect_file` does from a file), and the bytes cross the
worker pipe. Two things the local path never met: a blank Google Doc exports
as zero bytes, which is content and not "no input"; and a Doc exported as
Markdown carries base64 images and backslash escapes that would otherwise be
indexed as text.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.documents.extract.detect import detect_bytes, sniff
from app.documents.extract.dispatch import extract_any
from app.documents.extract.formats.markdown import strip_data_images, unescape_markdown
from app.documents.extract.pool import ExtractorPool
from app.documents.types import (
    EXTRACT_METADATA_ONLY,
    EXTRACT_OK,
    EXTRACT_PARTIAL,
    KIND_ARCHIVE,
    KIND_DOC,
    KIND_DOCX,
    KIND_IMAGE,
    KIND_MARKDOWN,
    KIND_PPTX,
    KIND_TEXT,
    KIND_XLS,
    KIND_XLSX,
    ExtractRequest,
)

from .extract_fixtures import make_cfb, make_docx, make_word_streams, zip_bytes


@pytest.fixture(scope="module")
def pool(tmp_path_factory):
    p = ExtractorPool(1, tmp_root=str(tmp_path_factory.mktemp("extract-bytes")), mem_limit_mb=512)
    yield p
    p.close()


def _xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Budget"
    wb.active.append(["item", "cost"])
    wb.active.append(["robot arm", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly plan"
    slide.placeholders[1].text = "Ship the tracker"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _late_directory_xlsx() -> bytes:
    """An xlsx whose first member is 20 KB of noise, so nothing decisive is in
    the first 8 KiB: only the whole directory names the kind."""
    ct = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.'
          'spreadsheetml.sheet.main+xml"/></Types>')
    noise = bytes((i * 131 + 7) % 251 for i in range(20_000))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        zf.writestr("docProps/thumbnail.bin", noise)
        zf.writestr("[Content_Types].xml", ct)
        zf.writestr("xl/workbook.xml", "<workbook/>")
    return buf.getvalue()


# ── Empty content ──────────────────────────────────────────────────────────


@pytest.mark.parametrize("name,kind", [("Untitled.md", "markdown"), ("Untitled.txt", "text")])
def test_empty_bytes_are_empty_content_through_a_real_worker(pool, name, kind) -> None:
    r = pool.extract(ExtractRequest(name=name, kind=kind, data=b""), timeout_s=60)
    assert (r.status, r.reason, r.blocks) == (EXTRACT_OK, None, [])
    # In process it was already right; the pipe used to turn b"" into None.
    assert extract_any(ExtractRequest(name=name, kind=kind, data=b"")).status == EXTRACT_OK


def test_empty_bytes_detect_by_name() -> None:
    assert detect_bytes(b"", "Untitled.md")[0] == KIND_MARKDOWN
    assert detect_bytes(b"", "Untitled")[0] == KIND_TEXT


# ── Detection from bytes ───────────────────────────────────────────────────


def test_detect_bytes_names_office_files_without_an_extension() -> None:
    assert detect_bytes(make_docx([(None, "x")]), "Report")[0] == KIND_DOCX
    assert detect_bytes(_xlsx(), "Budget")[0] == KIND_XLSX
    assert detect_bytes(_pptx(), "Deck")[0] == KIND_PPTX
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
    assert detect_bytes(png, "drawing.png") == (KIND_IMAGE, "image/png")


def test_detect_bytes_reads_the_whole_zip_directory() -> None:
    data = _late_directory_xlsx()
    # The head alone cannot tell, and there is no extension to break the tie...
    assert sniff(data[:8192], data[-512:], "")[0] == KIND_ARCHIVE
    # ...but the directory at the end of the bytes can.
    assert detect_bytes(data, "Budget")[0] == KIND_XLSX
    # And a plain zip named .docx is still an archive once its directory is read.
    assert detect_bytes(zip_bytes({"notes.txt": "x"}), "fake.docx") == (KIND_ARCHIVE, "application/zip")


def test_detect_bytes_reads_ole_directories() -> None:
    assert detect_bytes(make_cfb(make_word_streams([("hi\r", True)])), "letter") == (KIND_DOC, "application/msword")
    assert detect_bytes(make_cfb({"Workbook": b"x"}), "book")[0] == KIND_XLS


def test_ole_bytes_reach_olefile_as_a_stream(monkeypatch) -> None:
    # olefile reads bytes shorter than 1536 as a *file name*; detect_bytes
    # must hand it a stream, and junk after the magic falls back to the head.
    import olefile

    seen: list[type] = []
    real = olefile.OleFileIO

    def spy(source, *args, **kwargs):
        seen.append(type(source))
        return real(source, *args, **kwargs)

    monkeypatch.setattr(olefile, "OleFileIO", spy)
    junk = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"x" * 40
    assert detect_bytes(junk, "old.doc")[0] == KIND_DOC  # the name breaks the tie
    assert detect_bytes(make_cfb({"Workbook": b"x"}), "b")[0] == KIND_XLS
    assert seen and bytes not in seen


# ── Office files and archives through a real worker ────────────────────────


@pytest.mark.parametrize("name,maker,status", [
    ("Report", lambda: make_docx([("Heading1", "Minutes"), (None, "The robot follows a red ball.")]), EXTRACT_OK),
    ("Budget", _xlsx, EXTRACT_OK),
    ("Deck", _pptx, EXTRACT_OK),
    ("letter", lambda: make_cfb(make_word_streams([("Old Word text\r", True)])), EXTRACT_OK),
    ("bundle.zip", lambda: zip_bytes({"inner.txt": "x", "b/c.txt": "y"}), EXTRACT_METADATA_ONLY),
])
def test_office_files_from_bytes_through_a_real_worker(pool, name, maker, status) -> None:
    data = maker()
    kind = detect_bytes(data, name)[0]
    r = pool.extract(ExtractRequest(name=name, kind=kind, data=data), timeout_s=120)
    assert r.status == status, (name, kind, r.reason, r.doc_meta)
    assert r.kind == kind
    if status == EXTRACT_OK:
        assert r.blocks
    else:
        assert sorted(r.listing) == ["b/c.txt", "inner.txt"]


def test_xlsx_and_pptx_keep_their_locators(pool) -> None:
    r = pool.extract(ExtractRequest(name="Budget.xlsx", kind=KIND_XLSX, data=_xlsx()), timeout_s=120)
    assert any(b.locator.get("sheet") == "Budget" for b in r.blocks)
    r = pool.extract(ExtractRequest(name="Deck.pptx", kind=KIND_PPTX, data=_pptx()), timeout_s=120)
    assert any(b.locator.get("slide") == 1 and "Quarterly plan" in b.text for b in r.blocks)


# ── A Google Doc exported as Markdown ──────────────────────────────────────

_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk" * 1500  # ~90 KB
EXPORT = f"""# Luật Đất đai

## Điều 12\\. Phạm vi điều chỉnh

Luật này quy định về chế độ sở hữu đất đai \\(xem Điều 13\\)\\.

![][image1]

\\# not a heading, \\- not a list

1\\. still one paragraph

Inline `a\\.b` code keeps its escape.

```
path\\.with\\.escapes
```

| Cột | Giá trị |
|---|---|
| a\\_b | 1\\.5 |

[image1]: <data:image/png;base64,{_B64}>
"""
_EXPORT_LIMITS = {"md_locators": "structure", "md_export": True}


def _check_export(r) -> None:
    assert r.status == EXTRACT_OK, (r.reason, r.doc_meta)
    texts = [b.text for b in r.blocks]
    joined = "\n".join(texts)
    assert "data:image" not in joined and "base64" not in joined and "[image1]" not in joined
    assert r.doc_meta["title"] == "Luật Đất đai"
    heading = next(b for b in r.blocks if b.role == "heading" and b.level == 2)
    assert heading.text == "Điều 12. Phạm vi điều chỉnh"
    assert "Luật này quy định về chế độ sở hữu đất đai (xem Điều 13)." in texts
    # Escapes kept the structure while parsing: neither line became a heading or list.
    assert "# not a heading, - not a list" in texts
    assert "1. still one paragraph" in texts
    assert not any(b.role in ("heading", "list") and "not a" in b.text for b in r.blocks)
    # Code keeps its backslashes, fenced or inline.
    assert "Inline `a\\.b` code keeps its escape." in texts
    assert any(b.role == "code" and "path\\.with\\.escapes" in b.text for b in r.blocks)
    table = next(b for b in r.blocks if b.role == "table")
    assert "| a_b | 1.5 |" in table.text
    # Structure locators: the heading path and paragraph numbers, never lines
    # of an export the user cannot open.
    for block in r.blocks:
        assert "line_start" not in block.locator
        assert "para" in block.locator and "heading" in block.locator
    para = next(b for b in r.blocks if b.text.startswith("Luật này"))
    assert para.locator["heading"] == ["Luật Đất đai", "Điều 12. Phạm vi điều chỉnh"]


def test_markdown_export_through_a_real_worker(pool) -> None:
    r = pool.extract(ExtractRequest(name="Luật.md", kind=KIND_MARKDOWN, data=EXPORT.encode("utf-8"),
                                    limits=_EXPORT_LIMITS), timeout_s=120)
    _check_export(r)


def test_markdown_export_in_process() -> None:
    _check_export(extract_any(ExtractRequest(name="Luật.md", kind=KIND_MARKDOWN,
                                             data=EXPORT.encode("utf-8"), limits=dict(_EXPORT_LIMITS))))


def test_a_plain_markdown_file_is_untouched_by_the_export_rules() -> None:
    r = extract_any(ExtractRequest(name="notes.md", kind=KIND_MARKDOWN, data=b"# A\n\nx \\. y\n"))
    assert [b.text for b in r.blocks] == ["A", "x \\. y"]
    assert r.blocks[1].locator == {"line_start": 3, "line_end": 3}


def test_images_are_stripped_before_the_text_budget_applies() -> None:
    # ~90 KB of base64 against a 50 KB budget: the text itself fits.
    limits = {"max_text_mb": 0.05}
    data = EXPORT.encode("utf-8")
    assert len(data) > 0.05 * 1024 * 1024
    plain = extract_any(ExtractRequest(name="d.md", kind=KIND_MARKDOWN, data=data, limits=dict(limits)))
    assert plain.status == EXTRACT_PARTIAL
    export = extract_any(ExtractRequest(name="d.md", kind=KIND_MARKDOWN, data=data,
                                        limits={**limits, **_EXPORT_LIMITS}))
    assert export.status == EXTRACT_OK


def test_strip_data_images_keeps_ordinary_links_and_alt_text() -> None:
    text = ("See ![chart][image2] and ![logo](data:image/gif;base64,R0lGOD) and [site][ref].\n"
            "![kept][other]\n"
            "[image2]: <data:image/jpeg;base64,/9j/4AAQ>\n"
            "[ref]: https://example.com\n"
            "[other]: https://example.com/p.png\n")
    out = strip_data_images(text)
    assert out.splitlines() == [
        "See chart and logo and [site][ref].",
        "![kept][other]",
        "[ref]: https://example.com",
        "[other]: https://example.com/p.png",
    ]


def test_unescape_markdown() -> None:
    assert unescape_markdown(r"Điều 12\. a\-b \*c\* \_d\_ \#e \+f \!g \(h\) \[i\]") \
        == "Điều 12. a-b *c* _d_ #e +f !g (h) [i]"
    assert unescape_markdown(r"keep \| pipes and \\ and `co\.de`") == r"keep \| pipes and \\ and `co\.de`"
    assert unescape_markdown("no escapes") == "no escapes"


def test_the_worker_frees_the_request_before_parsing(pool) -> None:
    # Not a memory measurement: a larger in-memory file still round-trips
    # after the frame and message copies are dropped early.
    body = ("Dòng văn bản số {}.\n\n" * 20_000).format(*range(20_000)).encode("utf-8")
    r = pool.extract(ExtractRequest(name="big.txt", kind=KIND_TEXT, data=body), timeout_s=120)
    assert r.status == EXTRACT_OK
    assert r.blocks[0].text == "Dòng văn bản số 0." and r.blocks[-1].text == "Dòng văn bản số 19999."
