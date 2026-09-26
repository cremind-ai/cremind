"""Per-format extraction, in process through ``extract_any``.

Every fixture is generated here or in :mod:`.extract_fixtures`; the real
parser libraries run against them.
"""

from __future__ import annotations

import base64
import datetime
import gzip
import hashlib
import io
import struct
import sys
import tarfile
import unicodedata
import zipfile
from email.message import EmailMessage

import pytest

from app.documents.extract.dispatch import DEFAULT_LIMITS, extract_any
from app.documents.extract.formats._base import group_rows
from app.documents.extract.formats.legacy import doc_sections, parse_piece_table, piece_text
from app.documents.extract.formats.rtf import rtf_to_text
from app.documents.types import (
    ANCHOR_HARD,
    ANCHOR_NONE,
    ANCHOR_SOFT,
    EXTRACT_ERROR,
    EXTRACT_METADATA_ONLY,
    EXTRACT_OK,
    EXTRACT_PARTIAL,
    ExtractRequest,
    ExtractResult,
)

from .extract_fixtures import (
    encrypt_pdf,
    make_cfb,
    make_docx,
    make_epub,
    make_odf,
    make_pdf,
    make_ppt_streams,
    make_word_streams,
)


def run(name: str, kind: str, data: bytes | str, **limits) -> ExtractResult:
    raw = data.encode("utf-8") if isinstance(data, str) else data
    result = extract_any(ExtractRequest(name=name, kind=kind, data=raw, limits=limits))
    # Whatever came back must survive the worker's JSON pipe unchanged.
    import json

    assert ExtractResult.from_dict(json.loads(json.dumps(result.to_dict()))).to_dict() == result.to_dict()
    return result


def texts(result: ExtractResult) -> list[str]:
    return [b.text for b in result.blocks]


def test_default_limits() -> None:
    assert DEFAULT_LIMITS["max_pages"] == 2000
    assert DEFAULT_LIMITS["max_text_mb"] == 20
    assert DEFAULT_LIMITS["max_rows_per_sheet"] == 5000
    assert DEFAULT_LIMITS["max_archive_names"] == 500
    assert DEFAULT_LIMITS["max_ocr_pages"] == 50


# ── Markdown ───────────────────────────────────────────────────────────────

MARKDOWN = """---
title: Robot notes
author: Lee
---
# Robot

Intro line one
line two

## Tracking

```bash
# not a heading

echo hi
```

#### Detail

- item a
- item b

> quoted

Setext Title
============
"""


def test_markdown_headings_fences_and_line_locators() -> None:
    r = run("notes.md", "markdown", MARKDOWN)
    assert r.status == EXTRACT_OK
    assert r.doc_meta["title"] == "Robot notes"
    assert r.doc_meta["author"] == "Lee"
    by_text = {b.text: b for b in r.blocks}
    assert by_text["Robot"].anchor == ANCHOR_HARD and by_text["Robot"].level == 1
    assert by_text["Tracking"].anchor == ANCHOR_HARD and by_text["Tracking"].role == "heading"
    assert by_text["Detail"].anchor == ANCHOR_SOFT and by_text["Detail"].level == 4
    assert by_text["Setext Title"].level == 1
    fence = next(b for b in r.blocks if b.role == "code")
    assert "# not a heading" in fence.text and "echo hi" in fence.text  # one block across the blank line
    assert not any(b.text == "not a heading" for b in r.blocks)
    intro = by_text["Intro line one\nline two"]
    assert intro.locator == {"line_start": 7, "line_end": 8}
    assert by_text["- item a\n- item b"].role == "list"
    assert by_text["> quoted"].role == "quote"


def test_markdown_large_table_is_split_with_header_repeated() -> None:
    rows = "\n".join(f"| r{i} | {i * 7} |" for i in range(1, 121))
    r = run("t.md", "markdown", f"# T\n\n| name | value |\n|---|---|\n{rows}\n")
    tables = [b for b in r.blocks if b.role == "table"]
    assert len(tables) >= 3
    for block in tables:
        lines = block.text.split("\n")
        assert lines[0] == "| name | value |" and lines[1] == "|---|---|"
        assert len(lines) - 2 <= 40
    covered = [n for b in tables for n in range(b.locator["rows"][0], b.locator["rows"][1] + 1)]
    assert covered == list(range(1, 121))


def test_markdown_unclosed_fence_swallows_the_rest() -> None:
    r = run("x.md", "markdown", "# A\n\n```\n# B\ntext\n")
    assert [b.role for b in r.blocks] == ["heading", "code"]


# ── Plain text, JSON, code ─────────────────────────────────────────────────


def test_text_paragraphs_with_line_locators() -> None:
    r = run("a.txt", "text", "first para\nstill first\n\n\nsecond para\n")
    assert texts(r) == ["first para\nstill first", "second para"]
    assert [b.locator for b in r.blocks] == [{"line_start": 1, "line_end": 2}, {"line_start": 5, "line_end": 5}]
    assert r.doc_meta["encoding"] == "utf-8"


_TONES = {chr(0x300), chr(0x301), chr(0x303), chr(0x309), chr(0x323)}


def _cp1258(text: str) -> bytes:
    """Vietnamese the way Windows-1258 stores it: precomposed base letters
    (ô, ơ, ư, ă, â, ê, đ) followed by a combining tone mark."""
    out = []
    for ch in text:
        parts = unicodedata.normalize("NFD", ch)
        tones = "".join(c for c in parts if c in _TONES)
        out.append(unicodedata.normalize("NFC", "".join(c for c in parts if c not in _TONES)) + tones)
    return "".join(out).encode("cp1258")


def test_text_charset_detection_and_bom() -> None:
    sentence = "Công văn số 12 về việc thực hiện hợp đồng thuê nhà. "
    r = run("cv.txt", "text", _cp1258(sentence * 20))
    assert r.doc_meta["encoding"] == "cp1258"
    # Decoded as base + combining marks, composed back to what people type.
    assert r.blocks[0].text.startswith(sentence.strip())
    ru = run("ru.txt", "text", ("Привет, как дела? Это тест кодировки. " * 20).encode("cp1251"))
    assert ru.blocks[0].text.startswith("Привет, как дела?")
    r16 = run("u16.txt", "text", "Xin chào thế giới\n\nđoạn hai".encode("utf-16"))
    assert texts(r16) == ["Xin chào thế giới", "đoạn hai"]


def test_text_budget_marks_partial_and_keeps_the_head() -> None:
    body = "\n\n".join(f"paragraph number {i} " * 5 for i in range(2000))
    r = run("big.txt", "text", body, max_text_mb=0.01)
    assert r.status == EXTRACT_PARTIAL and r.reason == "too_large"
    assert r.blocks and r.blocks[0].text.startswith("paragraph number 0")
    assert sum(len(t) for t in texts(r)) <= 0.01 * 1024 * 1024


def test_minified_json_is_pretty_printed() -> None:
    r = run("d.json", "json", '{"project":"robot","tags":["vision","motion"],"n":3}')
    assert '"project": "robot"' in r.blocks[0].text
    assert "line_start" not in r.blocks[0].locator  # the indented text has no lines on disk


PYTHON = '''"""Module doc."""
import os
import sys as system
from app.documents import types
if True:
    from json import loads

CONSTANT = 1


# Explains helper.
@decorator
def helper(x):
    return x


class Tracker:
    def a(self):
        return 1
'''


def test_python_blocks_follow_definitions() -> None:
    r = run("track.py", "code", PYTHON)
    meta = r.blocks[0]
    assert meta.role == "meta" and meta.text == "imports: os, sys, app.documents, json"
    helper = next(b for b in r.blocks if "def helper" in b.text)
    assert helper.anchor == ANCHOR_SOFT
    assert helper.text.startswith("# Explains helper.\n@decorator")
    assert helper.locator == {"line_start": 11, "line_end": 14}
    tracker = next(b for b in r.blocks if b.text.startswith("class Tracker"))
    assert tracker.anchor == ANCHOR_SOFT and tracker.locator == {"line_start": 17, "line_end": 19}
    constant = next(b for b in r.blocks if b.text.startswith("CONSTANT"))
    assert constant.anchor == ANCHOR_NONE


def test_long_python_class_splits_at_methods() -> None:
    methods = "\n".join(f"    def m{i}(self):\n" + "".join(f"        x{j} = {j}\n" for j in range(8))
                        for i in range(10))
    r = run("big.py", "code", f"class Big:\n    '''doc'''\n\n{methods}")
    pieces = [b for b in r.blocks if b.role == "code"]
    assert pieces[0].anchor == ANCHOR_SOFT and pieces[0].text.startswith("class Big")
    assert all(b.anchor == ANCHOR_NONE for b in pieces[1:])
    assert any(b.text.lstrip().startswith("def m3") for b in pieces)


def test_non_python_code_uses_regex_anchors() -> None:
    js = ("import x from 'y';\n\n/** Adds. */\nexport function add(a, b) {\n  return a + b;\n}\n\n"
          "class Robot {\n  move() {}\n}\n")
    r = run("robot.js", "code", js)
    add = next(b for b in r.blocks if "function add" in b.text)
    assert add.anchor == ANCHOR_SOFT and add.text.startswith("/** Adds. */")
    robot = next(b for b in r.blocks if b.text.startswith("class Robot"))
    assert robot.anchor == ANCHOR_SOFT
    assert r.blocks[0].text == "import x from 'y';" and r.blocks[0].anchor == ANCHOR_NONE


def test_invalid_python_falls_back_to_regex() -> None:
    r = run("py2.py", "code", "print 'hello'\n\ndef f():\n    pass\n")
    assert any(b.anchor == ANCHOR_SOFT and b.text.startswith("def f") for b in r.blocks)


# ── CSV and spreadsheets ───────────────────────────────────────────────────


def _csv(n: int, insert_at: int | None = None) -> str:
    rows = [f"item{i},{i * 3},note {i}" for i in range(n)]
    if insert_at is not None:
        rows.insert(insert_at, "inserted,0,new row")
    return "name,qty,note\n" + "\n".join(rows) + "\n"


def test_csv_groups_repeat_the_header() -> None:
    r = run("inv.csv", "csv", _csv(300))
    assert r.status == EXTRACT_OK
    first_data = 1
    for block in r.blocks:
        lines = block.text.split("\n")
        assert lines[0] == "| name | qty | note |"
        assert 1 <= len(lines) - 2 <= 40
        assert block.locator["rows"][0] == first_data
        first_data = block.locator["rows"][1] + 1
        assert block.role == "table"
    assert first_data == 301


def test_csv_groups_resynchronise_after_an_insert() -> None:
    """One inserted row near the top changes a couple of blocks, not all
    that follow it (fixed 40-row groups would change every later block)."""
    before = set(texts(run("a.csv", "csv", _csv(400))))
    after = texts(run("a.csv", "csv", _csv(400, insert_at=5)))
    changed = [t for t in after if t not in before]
    assert len(after) > 8
    assert len(changed) <= 2


def test_group_rows_bounds() -> None:
    rows = [(i, [f"value {i}"]) for i in range(1000)]
    groups = list(group_rows(rows))
    assert all(1 <= len(g) <= 40 for g in groups)
    assert [r for g in groups for r in g] == rows


def test_tsv_and_row_cap() -> None:
    data = "a\tb\n" + "\n".join(f"{i}\tx" for i in range(30))
    r = run("t.tsv", "csv", data, max_rows_per_sheet=10)
    assert r.status == EXTRACT_PARTIAL and r.reason == "too_large"
    assert "| a | b |" in r.blocks[0].text
    assert r.blocks[-1].locator["rows"][1] == 10


def test_xlsx_sheets_ranges_and_cap() -> None:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Doanh thu"
    ws.append(["Tháng", "Số tiền", "Ngày"])
    for i in range(90):
        ws.append([f"T{i}", i * 1.5, datetime.date(2024, 1, 1) + datetime.timedelta(days=i)])
    other = wb.create_sheet("Chi phí")
    other.append(["Mục", "Tiền"])
    other.append(["Điện", 100])
    buf = io.BytesIO()
    wb.save(buf)
    r = run("book.xlsx", "xlsx", buf.getvalue(), max_rows_per_sheet=60)
    assert r.status == EXTRACT_PARTIAL and r.reason == "too_large"
    headings = [b for b in r.blocks if b.role == "heading"]
    assert [h.text for h in headings] == ["Doanh thu", "Chi phí"]
    assert all(h.anchor == ANCHOR_HARD for h in headings)
    tables = [b for b in r.blocks if b.role == "table" and b.locator["sheet"] == "Doanh thu"]
    assert tables[0].locator["range"].startswith("A2:C")
    assert tables[-1].locator["range"].endswith("61")  # 60 data rows after the header row
    assert "| Tháng | Số tiền | Ngày |" in tables[0].text
    assert "2024-01-01" in tables[0].text and "| 1.5 |" in tables[0].text
    assert r.doc_meta["sheets"] == 2


# ── PDF ────────────────────────────────────────────────────────────────────

_PDF_PAGES = [
    [(26, 72, 720, "Annual Report"), (11, 72, 690, "The first body line"), (11, 72, 676, "continues here"),
     (16, 72, 640, "Revenue"), (11, 72, 620, "Revenue grew a lot"), (14, 72, 590, "Minor note heading"),
     (11, 72, 570, "small body")],
    [(11, 72, 720, "Second page body")],
]


def test_pdf_headings_pages_and_metadata() -> None:
    pdf = make_pdf(_PDF_PAGES, info={"Title": "Report 2023", "Author": "Ann",
                                     "CreationDate": "D:20230115103000+07'00'", "ModDate": "D:2023021608"})
    r = run("r.pdf", "pdf", pdf)
    assert r.status == EXTRACT_OK
    assert r.doc_meta["title"] == "Report 2023" and r.doc_meta["author"] == "Ann"
    assert r.doc_meta["created"] == "2023-01-15T10:30:00+07:00"
    assert r.doc_meta["modified"] == "2023-02-16T08:00:00"
    assert r.doc_meta["pages"] == 2
    by_text = {b.text: b for b in r.blocks}
    assert by_text["Annual Report"].anchor == ANCHOR_HARD and by_text["Annual Report"].level == 1
    assert by_text["Revenue"].anchor == ANCHOR_HARD and by_text["Revenue"].level == 2
    assert by_text["Minor note heading"].anchor == ANCHOR_SOFT and by_text["Minor note heading"].level == 3
    body = by_text["The first body line\ncontinues here"]
    assert body.locator == {"page": 1, "line_start": 2, "line_end": 3}
    page2 = by_text["Second page body"]
    assert page2.anchor == ANCHOR_SOFT and page2.locator["page"] == 2  # the page break


def test_pdf_running_headers_are_dropped() -> None:
    bodies = ["Alpha findings", "Beta discussion", "Gamma results", "Delta appendix"]
    pages = [[(11, 72, 760, "ACME Confidential"), (11, 72, 700, body),
              (11, 72, 40, f"Page {n} of 4")] for n, body in enumerate(bodies, start=1)]
    r = run("h.pdf", "pdf", make_pdf(pages))
    assert texts(r) == bodies


def test_pdf_scanned_page_is_rendered_for_ocr() -> None:
    r = run("scan.pdf", "pdf", make_pdf([[(11, 72, 700, "A text page")], []], scanned=frozenset({2})))
    assert r.doc_meta["scanned_pages"] == [2]
    assert len(r.ocr_pages) == 1
    page = r.ocr_pages[0]
    png = base64.b64decode(page["png_b64"])
    assert png.startswith(b"\x89PNG") and page["page"] == 2
    assert page["sha256"] == hashlib.sha256(png).hexdigest()
    from PIL import Image

    with Image.open(io.BytesIO(png)) as im:
        assert im.size == (1224, 1584)  # 612x792 pt at 144 dpi


def test_pdf_ocr_cap_marks_partial() -> None:
    pdf = make_pdf([[], [], []], scanned=frozenset({1, 2, 3}))
    r = run("scan.pdf", "pdf", pdf, max_ocr_pages=2)
    assert len(r.ocr_pages) == 2 and r.status == EXTRACT_PARTIAL


def test_pdf_page_cap() -> None:
    r = run("p.pdf", "pdf", make_pdf([[(11, 72, 700, f"page {n}")] for n in range(1, 6)]), max_pages=2)
    assert r.status == EXTRACT_PARTIAL and r.reason == "too_large"
    assert texts(r) == ["page 1", "page 2"]
    assert r.doc_meta["pages"] == 5


def test_pdf_damaged_page_keeps_the_others(monkeypatch) -> None:
    from app.documents.extract.formats import pdf as pdf_mod

    real = pdf_mod._page_lines

    def flaky(page):
        if page.page_number == 2:
            raise ValueError("bad content stream")
        return real(page)

    monkeypatch.setattr(pdf_mod, "_page_lines", flaky)
    r = run("p.pdf", "pdf", make_pdf([[(11, 72, 700, f"page {n}")] for n in range(1, 4)]))
    assert (r.status, r.reason) == (EXTRACT_PARTIAL, "corrupt")
    assert texts(r) == ["page 1", "page 3"]
    assert r.doc_meta["unreadable_pages"] == [2]


def test_pdf_encryption() -> None:
    pytest.importorskip("pypdf")
    plain = make_pdf([[(11, 72, 700, "secret body")]])
    locked = run("l.pdf", "pdf", encrypt_pdf(plain, "user-pass"))
    assert locked.status == EXTRACT_METADATA_ONLY and locked.reason == "encrypted"
    assert not locked.blocks
    owner_only = run("o.pdf", "pdf", encrypt_pdf(plain, ""))
    assert owner_only.status == EXTRACT_OK and texts(owner_only) == ["secret body"]


def test_corrupt_pdf_is_an_error() -> None:
    r = run("bad.pdf", "pdf", b"%PDF-1.4\nthis is not really a pdf")
    assert r.status == EXTRACT_ERROR and r.reason == "corrupt"
    assert "error" in r.doc_meta


# ── Office (markitdown) ────────────────────────────────────────────────────


def test_docx_structure_locators_and_core_meta() -> None:
    pytest.importorskip("mammoth")
    docx = make_docx([("Heading1", "Chương I"), (None, "Điều khoản chung"), (None, "Đoạn hai"),
                      ("Heading2", "Mục 1"), (None, "Nội dung mục")])
    r = run("hd.docx", "docx", docx)
    assert r.status == EXTRACT_OK
    assert r.doc_meta["title"] == "Quarterly Report"
    assert r.doc_meta["author"] == "Alice Nguyen" and r.doc_meta["last_modified_by"] == "Bob Tran"
    assert r.doc_meta["created"] == "2024-01-02T03:04:05+00:00"
    assert r.doc_meta["pages"] == 3 and r.doc_meta["words"] == 120
    by_text = {b.text: b for b in r.blocks}
    assert by_text["Chương I"].anchor == ANCHOR_HARD
    section = by_text["Nội dung mục"]
    assert section.locator == {"heading": ["Chương I", "Mục 1"], "para": [5, 5]}
    assert "line_start" not in section.locator


def test_html_via_markitdown() -> None:
    pytest.importorskip("bs4")
    html = ("<html><head><title>Trang</title><meta charset='utf-8'></head><body><h1>Tiêu đề</h1>"
            "<p>Xin chào</p><script>var hidden = 1;</script></body></html>")
    r = run("p.html", "html", html)
    assert r.doc_meta["title"] == "Trang"
    assert texts(r) == ["Tiêu đề", "Xin chào"]


def test_html_legacy_charset_from_meta() -> None:
    pytest.importorskip("bs4")
    html = '<html><head><meta charset="windows-1258"></head><body><p>Vi\xeat Nam</p></body></html>'
    r = run("old.htm", "html", html.encode("latin-1"))
    assert texts(r) == ["Viêt Nam"]  # 0xEA is ê in cp1258, not the UTF-8 the converter would assume


def test_epub_via_markitdown() -> None:
    r = run("b.epub", "epub", make_epub())
    assert r.doc_meta["title"] == "The Robot Book"
    heading = next(b for b in r.blocks if b.role == "heading")
    assert heading.text == "Motion tracking" and heading.anchor == ANCHOR_HARD
    assert "The robot follows a red ball." in texts(r)


def test_msg_via_markitdown() -> None:
    pytest.importorskip("olefile")
    msg = make_cfb({
        "__substg1.0_0037001F": "Hợp đồng thuê nhà".encode("utf-16-le"),
        "__substg1.0_0C1F001F": "lan@example.vn".encode("utf-16-le"),
        "__substg1.0_1000001F": "Chào anh,\r\n\r\nGửi hợp đồng.".encode("utf-16-le"),
    })
    r = run("m.msg", "msg", msg)
    assert r.doc_meta["title"] == "Hợp đồng thuê nhà" and r.doc_meta["author"] == "lan@example.vn"
    assert "Gửi hợp đồng." in texts(r)
    assert not any("\x00" in t for t in texts(r))  # the zero padding is cleaned


def test_pptx_one_hard_block_per_slide() -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Robot demo"
    s1.placeholders[1].text = "Tracks motion"
    s1.notes_slide.notes_text_frame.text = "Mention the camera"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Results"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (row, col), text in {(0, 0): "metric", (0, 1): "value", (1, 0): "speed", (1, 1): "3"}.items():
        table.cell(row, col).text = text
    buf = io.BytesIO()
    prs.save(buf)
    r = run("deck.pptx", "pptx", buf.getvalue())
    assert texts(r) == ["Robot demo\nTracks motion\nNotes: Mention the camera", "Results\nmetric | value\nspeed | 3"]
    assert [b.locator for b in r.blocks] == [{"slide": 1}, {"slide": 2}]
    assert all(b.anchor == ANCHOR_HARD for b in r.blocks)
    assert r.doc_meta["slides"] == 2


# ── OpenDocument ───────────────────────────────────────────────────────────


def test_odt_text_headings_lists_and_meta() -> None:
    odt = make_odf("application/vnd.oasis.opendocument.text",
                   '<office:text><text:h text:outline-level="1">Điều 1</text:h>'
                   '<text:p>Phạm vi<text:s text:c="2"/>điều chỉnh<text:note><text:note-body>'
                   '<text:p>footnote</text:p></text:note-body></text:note></text:p>'
                   '<text:list><text:list-item><text:p>một</text:p></text:list-item>'
                   '<text:list-item><text:p>hai</text:p></text:list-item></text:list>'
                   '<text:table-of-content><text:p>TOC line</text:p></text:table-of-content></office:text>')
    r = run("a.odt", "odt", odt)
    assert texts(r) == ["Điều 1", "Phạm vi  điều chỉnh", "- một\n- hai"]  # text:s c=2; footnote and TOC out
    assert r.blocks[0].anchor == ANCHOR_HARD
    assert r.blocks[2].locator == {"heading": ["Điều 1"], "para": [3, 4]}
    assert r.doc_meta["author"] == "Lan" and r.doc_meta["last_modified_by"] == "Minh"
    assert r.doc_meta["created"] == "2023-05-01T10:00:00" and r.doc_meta["pages"] == 2


def test_ods_repeats_and_empty_tail() -> None:
    ods = make_odf("application/vnd.oasis.opendocument.spreadsheet",
                   '<office:spreadsheet><table:table table:name="Thu chi">'
                   '<table:table-row><table:table-cell><text:p>Ngày</text:p></table:table-cell>'
                   '<table:table-cell><text:p>Tiền</text:p></table:table-cell></table:table-row>'
                   '<table:table-row table:number-rows-repeated="3"><table:table-cell><text:p>x</text:p>'
                   '</table:table-cell><table:table-cell><text:p>5</text:p></table:table-cell></table:table-row>'
                   '<table:table-row table:number-rows-repeated="1048000">'
                   '<table:table-cell table:number-columns-repeated="1024"/></table:table-row>'
                   '</table:table></office:spreadsheet>')
    r = run("a.ods", "ods", ods)
    assert r.blocks[0].text == "Thu chi" and r.blocks[0].anchor == ANCHOR_HARD
    assert r.blocks[1].locator == {"sheet": "Thu chi", "range": "A2:B4"}
    assert r.blocks[1].text.count("| x | 5 |") == 3


def test_odp_slides_with_notes() -> None:
    odp = make_odf("application/vnd.oasis.opendocument.presentation",
                   '<office:presentation><draw:page><draw:frame><draw:text-box><text:p>Slide one</text:p>'
                   '</draw:text-box></draw:frame><presentation:notes><draw:frame><draw:text-box>'
                   '<text:p>note</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>'
                   '<draw:page><draw:frame><draw:text-box><text:p>Two</text:p></draw:text-box></draw:frame>'
                   '</draw:page></office:presentation>')
    r = run("a.odp", "odp", odp)
    assert texts(r) == ["Slide one\nNotes: note", "Two"]
    assert [b.locator for b in r.blocks] == [{"slide": 1}, {"slide": 2}]


# ── Legacy .doc / .ppt / .rtf ──────────────────────────────────────────────


def test_piece_table_parser_mixed_pieces() -> None:
    first, second = "Hello \x13 HYPERLINK \"u\" \x14world\x15\r", "Tiếng Việt\r"
    streams = make_word_streams([(first, True), (second, False)])
    word, table = streams["WordDocument"], streams["1Table"]
    fc_clx, lcb = struct.unpack_from("<II", word, 154 + 33 * 8)
    pieces = parse_piece_table(table[fc_clx:fc_clx + lcb])  # the leading Prc is skipped
    n1 = len(first)
    assert [(p[0], p[1], p[3]) for p in pieces] == [(0, n1, True), (n1, n1 + len(second), False)]
    # A range straddling the cp1252 piece and the UTF-16 piece.
    assert piece_text(word, pieces, n1 - 2, n1 + 5) == "\x15\rTiếng"
    sections = doc_sections(word, None, table)
    assert sections == [("main", ["Hello world", "Tiếng Việt"])]


def test_piece_table_rejects_garbage() -> None:
    from app.documents.extract.formats._base import LegacyFormatError

    with pytest.raises(LegacyFormatError):
        parse_piece_table(b"\x07\x00\x00")
    with pytest.raises(LegacyFormatError):
        parse_piece_table(b"\x02" + struct.pack("<I", 7) + b"\x00" * 7)


def test_doc_end_to_end_with_footnotes() -> None:
    pytest.importorskip("olefile")
    streams = make_word_streams([("Điều 1. Phạm vi\rNội dung\r", False), ("See note\r", True)],
                                footnote_chars=len("See note\r"))
    r = run("law.doc", "doc", make_cfb(streams))
    assert r.status == EXTRACT_OK
    assert texts(r) == ["Điều 1. Phạm vi", "Nội dung", "Footnotes", "See note"]
    assert r.blocks[1].locator == {"para": [2, 2]}


def test_doc_encrypted_and_old_versions() -> None:
    pytest.importorskip("olefile")
    enc = run("e.doc", "doc", make_cfb(make_word_streams([("x\r", True)], encrypted=True)))
    assert enc.status == EXTRACT_METADATA_ONLY and enc.reason == "encrypted"
    old = run("w6.doc", "doc", make_cfb(make_word_streams([("x\r", True)], nfib=0x0065)))
    assert old.status == EXTRACT_METADATA_ONLY and old.reason == "legacy_format"
    no_table = make_word_streams([("x\r", True)])
    del no_table["1Table"]
    missing = run("m.doc", "doc", make_cfb(no_table))
    assert missing.reason == "legacy_format"
    not_ole = run("n.doc", "doc", b"this is not an OLE file at all")
    assert not_ole.status == EXTRACT_ERROR and not_ole.reason == "corrupt"


def test_ppt_slides_in_order_with_text_boxes() -> None:
    pytest.importorskip("olefile")
    # Placeholders are UTF-16 TextChars atoms; the text boxes are 8-bit TextBytes atoms.
    ppt = make_cfb(make_ppt_streams([(["Giới thiệu", "Mục tiêu"], ["Café box"]), (["Kết quả"], ["box two"])]))
    r = run("deck.ppt", "ppt", ppt)
    assert texts(r) == ["Giới thiệu\nMục tiêu\nCafé box", "Kết quả\nbox two"]
    assert [b.locator for b in r.blocks] == [{"slide": 1}, {"slide": 2}]
    assert all(b.anchor == ANCHOR_HARD for b in r.blocks)


def test_ppt_without_persist_directory_uses_slide_list() -> None:
    pytest.importorskip("olefile")
    streams = make_ppt_streams([(["Only placeholders"], ["lost box"])], with_current_user=False)
    r = run("deck.ppt", "ppt", make_cfb(streams))
    assert texts(r) == ["Only placeholders"]


def test_ppt_garbage_stream_is_legacy_format() -> None:
    pytest.importorskip("olefile")
    r = run("x.ppt", "ppt", make_cfb({"PowerPoint Document": b"\x00" * 64}))
    assert r.status == EXTRACT_METADATA_ONLY and r.reason == "legacy_format"


# The RTF Unicode escape, spelled so no editor or tool turns it into a character.
U = b"\\u"

RTF = (rb"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\fswiss\fcharset0 Arial;}{\f1\fnil\fcharset163 VN Times;}}"
       rb"{\colortbl;\red0\green0\blue0;}{\stylesheet{\s1 heading 1;}}{\*\generator Msftedit;}"
       rb"{\info{\title H" + U + b"7907?p " + U + b"273?" + U + rb"7891?ng}{\author Lan}{\operator Minh}"
       rb"{\creatim\yr2021\mo3\dy4\hr5\min6}{\revtim\yr2022\mo1\dy2}}"
       rb"\uc1\pard Xin ch\'e0o {\f1 Vi\'eat} Nam\par "
       rb"Second {\*\bkmkstart b}para\line new line\tab end\par "
       rb"{\pict\wmetafile8 0102abcdef}{\field{\*\fldinst HYPERLINK \"x\"}{\fldrslt link}}\par "
       # A surrogate pair with "?" fallbacks, then \bin with four raw bytes that
       # would derail a tokenizer if they were not skipped.
       rb"\u-10179?\u-8704? emoji\bin4 {}\\ done\par}")


def test_rtf_stripper() -> None:
    paragraphs, meta = rtf_to_text(RTF)
    assert paragraphs[0] == "Xin chào Viêt Nam"   # \'e0 in cp1252, \'ea in the font's cp1258
    assert paragraphs[1] == "Second para\nnew line\tend"
    assert paragraphs[2] == "link"
    assert paragraphs[3] == chr(0xD83D) + chr(0xDE00) + " emoji done"  # a lone pair; clean_str joins it later
    assert meta == {"title": "Hợp đồng", "author": "Lan", "last_modified_by": "Minh",
                    "created": "2021-03-04T05:06:00", "modified": "2022-01-02T00:00:00"}


def test_rtf_through_dispatch_cleans_surrogates() -> None:
    r = run("a.rtf", "rtf", RTF)
    assert r.status == EXTRACT_OK
    assert texts(r)[3] == "\U0001f600 emoji done"
    assert r.blocks[0].locator == {"para": [1, 1]}
    assert r.doc_meta["title"] == "Hợp đồng"


# ── Email ──────────────────────────────────────────────────────────────────


def test_eml_parts() -> None:
    msg = EmailMessage()
    msg["From"] = "Lan <lan@example.vn>"
    msg["To"] = "an@example.vn"
    msg["Subject"] = "Báo cáo quý 3"
    msg["Date"] = "Tue, 01 Aug 2023 10:00:00 +0700"
    msg.set_content("Chào anh,\n\nĐây là báo cáo.\n")
    msg.add_alternative("<p>Chào anh (html)</p>", subtype="html")
    msg.add_attachment(b"%PDF-1.4", maintype="application", subtype="pdf", filename="bao-cao.pdf")
    r = run("m.eml", "eml", msg.as_bytes())
    header, body1, body2, attachments = r.blocks
    assert header.role == "meta" and header.anchor == ANCHOR_HARD and header.locator == {"part": "headers"}
    assert "Subject: Báo cáo quý 3" in header.text
    assert (body1.text, body1.anchor, body1.locator) == ("Chào anh,", ANCHOR_HARD, {"part": "body"})
    assert body2.anchor == ANCHOR_NONE
    assert attachments.text == "Attachments: bao-cao.pdf" and attachments.locator == {"part": "attachments"}
    assert r.doc_meta["created"] == "2023-08-01T10:00:00+07:00"
    assert r.doc_meta["attachments"] == ["bao-cao.pdf"]


def test_eml_html_only_body() -> None:
    msg = EmailMessage()
    msg["Subject"] = "html"
    msg.set_content("<html><body><h1>Hi</h1><p>one</p><script>x()</script><p>two</p></body></html>",
                    subtype="html")
    r = run("h.eml", "eml", msg.as_bytes())
    assert texts(r)[1:] == ["Hi", "one", "two"]


# ── Images ─────────────────────────────────────────────────────────────────


def _jpeg_with_exif(**overrides) -> bytes:
    from PIL import Image

    exif = Image.Exif()
    exif[0x010F] = "Apple"
    exif[0x0110] = "iPhone 15"
    exif[0x0112] = 6
    exif[0x0131] = "17.1"
    exif[0x8769] = {0x9003: overrides.get("taken", "2023:05:01 14:30:00"), 0x9011: "+07:00"}
    exif[0x8825] = {1: "N", 2: (21.0, 1.0, 30.0), 3: "E", 4: (105.0, 51.0, 0.0)}
    buf = io.BytesIO()
    Image.new("RGB", (640, 480), "green").save(buf, "JPEG", exif=exif.tobytes())
    return buf.getvalue()


def test_jpeg_exif() -> None:
    pytest.importorskip("PIL")
    r = run("dog.jpg", "image", _jpeg_with_exif())
    assert r.status == EXTRACT_OK and not r.blocks
    assert r.exif == {"taken_at": "2023-05-01T14:30:00+07:00", "make": "Apple", "model": "iPhone 15",
                      "software": "17.1", "orientation": 6, "gps": {"lat": 21.025, "lon": 105.85},
                      "width": 640, "height": 480}
    assert r.image == {"width": 640, "height": 480, "format": "JPEG", "is_camera_photo": True}


def test_zeroed_camera_clock_is_no_date() -> None:
    pytest.importorskip("PIL")
    r = run("x.jpg", "image", _jpeg_with_exif(taken="0000:00:00 00:00:00"))
    assert "taken_at" not in r.exif and r.image["is_camera_photo"] is False


def test_png_without_exif_is_not_a_camera_photo() -> None:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("L", (300, 200)).save(buf, "PNG")
    r = run("s.png", "image", buf.getvalue())
    assert r.exif == {"width": 300, "height": 200}
    assert r.image == {"width": 300, "height": 200, "format": "PNG", "is_camera_photo": False}


def test_heic_with_and_without_pillow_heif(monkeypatch) -> None:
    pillow_heif = pytest.importorskip("pillow_heif")
    from PIL import Image

    pillow_heif.register_heif_opener()
    exif = Image.Exif()
    exif[0x010F] = "Apple"
    exif[0x8769] = {0x9003: "2024:02:03 04:05:06"}
    buf = io.BytesIO()
    Image.new("RGB", (320, 240), "red").save(buf, format="HEIF", exif=exif.tobytes())
    r = run("IMG_0001.HEIC", "image", buf.getvalue())
    assert r.status == EXTRACT_OK and r.image["format"] == "HEIF" and r.image["is_camera_photo"] is True
    assert r.exif["taken_at"] == "2024-02-03T04:05:06"

    from app.documents.extract.formats import image as image_mod

    monkeypatch.setattr(image_mod, "_register_heif", lambda: False)
    missing = run("IMG_0001.HEIC", "image", buf.getvalue())
    assert missing.status == EXTRACT_METADATA_ONLY and missing.reason == "heic_unsupported"


# ── Archives, executables, metadata-only kinds ─────────────────────────────


def test_zip_listing_is_capped() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(7):
            zf.writestr(f"src/f{i}.py", "x")
    r = run("a.zip", "archive", buf.getvalue(), max_archive_names=5)
    assert r.status == EXTRACT_METADATA_ONLY and r.reason == "archive"
    assert r.listing == [f"src/f{i}.py" for i in range(5)]
    assert r.doc_meta["members"] == 7 and r.doc_meta["listing_truncated"] is True
    assert not r.blocks


def test_tar_gz_and_plain_gzip_listing() -> None:
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as tar:
        for name in ("proj/README.md", "proj/main.py"):
            info = tarfile.TarInfo(name)
            info.size = 2
            tar.addfile(info, io.BytesIO(b"hi"))
    r = run("p.tar.gz", "archive", buf.getvalue())
    assert r.listing == ["proj/README.md", "proj/main.py"]
    assert r.doc_meta["format"] == "tar" and r.doc_meta["compression"] == "gzip"

    gz = io.BytesIO()
    with gzip.GzipFile(filename="notes.txt", mode="wb", fileobj=gz, mtime=0) as fh:
        fh.write(b"hello")
    r = run("notes.txt.gz", "archive", gz.getvalue())
    assert r.listing == ["notes.txt"] and r.doc_meta["format"] == "gzip"


def test_unlistable_archive_is_still_metadata_only() -> None:
    r = run("x.7z", "archive", b"7z\xbc\xaf\x27\x1c" + b"\x00" * 64)
    assert (r.status, r.reason, r.listing, r.doc_meta["format"]) == (EXTRACT_METADATA_ONLY, "archive", [], "7z")


@pytest.mark.parametrize("head, expected", [
    (b"\x7fELF\x02\x01\x01" + b"\x00" * 9 + struct.pack("<HH", 3, 0xB7) + b"\x00" * 64,
     {"format": "elf", "arch": "arm64", "bits": 64, "type": "shared"}),
    (b"\x7fELF\x01\x01\x01" + b"\x00" * 9 + struct.pack("<HH", 2, 0x03) + b"\x00" * 64,
     {"format": "elf", "arch": "x86", "bits": 32, "type": "executable"}),
    (b"MZ" + b"\x00" * 58 + struct.pack("<I", 64) + b"PE\x00\x00" + struct.pack("<H", 0xAA64)
     + b"\x00" * 16 + struct.pack("<H", 0x2000) + b"\x00" * 64,
     {"format": "pe", "arch": "arm64", "dll": True}),
    (b"\xcf\xfa\xed\xfe" + struct.pack("<I", 0x01000007) + b"\x00" * 64,
     {"format": "macho", "arch": "x86_64", "bits": 64}),
    (b"\xca\xfe\xba\xbe" + struct.pack(">I", 2) + struct.pack(">5I", 0x01000007, 3, 0, 0, 0)
     + struct.pack(">5I", 0x0100000C, 0, 0, 0, 0),
     {"format": "macho", "arch": "universal", "archs": ["x86_64", "arm64"]}),
])
def test_executable_arch(head: bytes, expected: dict) -> None:
    r = run("bin", "executable", head)
    assert (r.status, r.reason) == (EXTRACT_METADATA_ONLY, "executable")
    assert r.doc_meta == expected


def test_dmg_and_real_interpreter() -> None:
    dmg = b"\x78\xda" + b"\x01" * 4000 + b"koly" + b"\x00" * 508
    assert run("a.dmg", "executable", dmg).doc_meta == {"format": "dmg"}
    r = extract_any(ExtractRequest(name="python", kind="executable", path=sys.executable))
    assert r.reason == "executable" and r.doc_meta["format"] in ("pe", "elf", "macho")
    assert r.doc_meta.get("arch")


@pytest.mark.parametrize("kind", ["audio", "video", "database", "font", "encrypted", "other"])
def test_metadata_only_kinds_use_the_kind_as_reason(kind: str) -> None:
    r = run("f", kind, b"\x00\x01\x02")
    assert (r.status, r.reason, r.blocks) == (EXTRACT_METADATA_ONLY, kind, [])


# ── Dispatcher failure handling ────────────────────────────────────────────


def test_missing_file_and_no_input(tmp_path) -> None:
    r = extract_any(ExtractRequest(name="gone.txt", kind="text", path=str(tmp_path / "gone.txt")))
    assert (r.status, r.reason) == (EXTRACT_ERROR, "not_found")
    r = extract_any(ExtractRequest(name="x", kind="text"))
    assert (r.status, r.reason) == (EXTRACT_ERROR, "no_input")


def test_missing_optional_dependency_is_awaiting_extractor(monkeypatch) -> None:
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "pdfplumber":
            raise ModuleNotFoundError("No module named 'pdfplumber'", name="pdfplumber")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    r = run("a.pdf", "pdf", make_pdf([[(11, 72, 700, "x")]]))
    assert (r.status, r.reason) == (EXTRACT_ERROR, "awaiting_extractor")
    assert r.doc_meta["missing"] == "pdfplumber"


def test_failure_after_blocks_keeps_them_as_partial(monkeypatch) -> None:
    from app.documents.extract.formats import text as text_mod

    def half(ctx) -> None:
        ctx.add("readable start")
        raise ValueError("parser blew up")

    monkeypatch.setattr(text_mod, "extract_text", half)
    r = run("a.txt", "text", "ignored")
    assert (r.status, r.reason) == (EXTRACT_PARTIAL, "corrupt")
    assert texts(r) == ["readable start"]
    assert r.doc_meta["error"] == "ValueError: parser blew up"


def test_malformed_limits_fall_back_to_defaults() -> None:
    r = run("a.txt", "text", "fine\n", max_text_mb="lots", max_pages=None)
    assert r.status == EXTRACT_OK and texts(r) == ["fine"]


def test_strings_are_cleaned_for_storage() -> None:
    r = run("n.txt", "text", "bad\x00byte and \x07bell\n\nok")
    assert texts(r) == ["badbyte and bell", "ok"]


def test_kind_is_detected_when_missing(tmp_path) -> None:
    path = tmp_path / "notes"
    path.write_bytes(make_pdf([[(11, 72, 700, "detected")]]))
    r = extract_any(ExtractRequest(name="notes", kind="", path=str(path)))
    assert r.kind == "pdf" and texts(r) == ["detected"]
