"""Tabular files: CSV/TSV, XLSX and legacy XLS.

Every format goes through the same shape: the first non-empty row is the
header, data rows are grouped by :func:`~._base.group_rows`, and each group
is rendered as a Markdown table with the header on top, so a chunk of rows
240-270 still says what its columns are. Each sheet opens with a hard-anchored
heading, so no chunk spans two sheets.

Legacy ``.xls`` is read with ``xlrd`` directly rather than through markitdown:
markitdown loads every sheet into pandas and emits one Markdown table per
sheet, which loses the sheet and cell coordinates a citation needs, and pulls
pandas into the worker for nothing.
"""

from __future__ import annotations

import csv
import io
from typing import Any, Iterable, Iterator

from app.documents.types import ANCHOR_HARD

from ._base import Ctx, EncryptedDocument, cell_text, group_rows, read_text, render_table, trim_row

# openpyxl pads missing rows up to the last one the sheet XML mentions; a
# formatted cell far down the sheet can mean a million empty rows. Past this
# many in a row the rest of the sheet is treated as empty.
_MAX_EMPTY_RUN = 10_000


def _column_letter(index: int) -> str:
    """1 → A, 27 → AA."""
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters or "A"


def _emit_sheet_rows(ctx: Ctx, rows: Iterable[tuple[int, list[str]]], *,
                     sheet: str | None, max_rows: int) -> None:
    """Header detection, the row cap and grouping for one sheet (or CSV).

    ``rows`` are ``(row_number, cells)`` with row numbers as the source counts
    them: spreadsheet rows for sheets, data-row numbers for CSV.
    """
    header: list[str] | None = None
    header_row = 0
    capped = False

    def data_rows() -> Iterator[tuple[int, list[str]]]:
        nonlocal header, header_row, capped
        count = 0
        for number, cells in rows:
            if header is None:
                header, header_row = cells, number
                continue
            count += 1
            if count > max_rows:
                capped = True
                return
            yield number, cells

    emitted = False
    for group in group_rows(data_rows()):
        emitted = True
        first, last = group[0][0], group[-1][0]
        width = max([len(header or [])] + [len(cells) for _, cells in group])
        if sheet is not None:
            locator: dict[str, Any] = {"sheet": sheet, "range": f"A{first}:{_column_letter(width)}{last}"}
        else:
            locator = {"rows": [first, last]}
        ctx.add(render_table(header, [cells for _, cells in group]), role="table", locator=locator)
    if header is not None and not emitted:
        # A sheet holding only its header row still says what it tracks.
        locator = {"sheet": sheet, "range": f"A{header_row}:{_column_letter(len(header))}{header_row}"} \
            if sheet is not None else {"rows": [0, 0]}
        ctx.add(render_table(header, []), role="table", locator=locator)
    if capped:
        ctx.mark_partial()


def _non_empty(rows: Iterable[tuple[int, list[str]]]) -> Iterator[tuple[int, list[str]]]:
    empty_run = 0
    for number, cells in rows:
        cells = trim_row(cells)
        if not cells:
            empty_run += 1
            if empty_run > _MAX_EMPTY_RUN:
                return
            continue
        empty_run = 0
        yield number, cells


# ── CSV / TSV ──────────────────────────────────────────────────────────────


def _csv_dialect(text: str, ext: str) -> str:
    if ext in (".tsv", ".tab"):
        return "\t"
    sample = text[:16384]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def extract_csv(ctx: Ctx) -> None:
    """CSV/TSV. Data rows are numbered from 1 after the header (the header is
    row 0), which is what ``locator.rows`` holds."""
    text = read_text(ctx)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_csv_dialect(text, ctx.ext))

    def numbered() -> Iterator[tuple[int, list[str]]]:
        # Data-row numbering starts after the header, which _emit_sheet_rows
        # takes from the first non-empty row.
        number = -1
        try:
            for raw in reader:
                cells = trim_row([cell_text(c) for c in raw])
                if not cells:
                    continue
                number += 1
                yield number, cells
        except csv.Error:
            # A malformed tail (a field over csv's size limit, a stray NUL):
            # keep the rows before it.
            ctx.mark_partial("corrupt")

    _emit_sheet_rows(ctx, numbered(), sheet=None, max_rows=ctx.limit("max_rows_per_sheet"))


# ── XLSX ───────────────────────────────────────────────────────────────────


def extract_xlsx(ctx: Ctx) -> None:
    import openpyxl

    from .office import core_properties

    ctx.result.doc_meta.update(core_properties(ctx))
    wb = openpyxl.load_workbook(ctx.source(), read_only=True, data_only=True)
    try:
        sheets = list(wb.worksheets)
        ctx.result.doc_meta["sheets"] = len(sheets)
        max_rows = ctx.limit("max_rows_per_sheet")
        for ws in sheets:
            # Stale <dimension> records are common (and make openpyxl pad every
            # row to the recorded width); read the rows as they are instead.
            ws.reset_dimensions()
            title = str(ws.title)
            ctx.add(title, anchor=ANCHOR_HARD, level=1, role="heading", locator={"sheet": title})
            rows = ((number, [cell_text(v) for v in row])
                    for number, row in enumerate(ws.iter_rows(values_only=True), start=1))
            _emit_sheet_rows(ctx, _non_empty(rows), sheet=title, max_rows=max_rows)
    finally:
        wb.close()


# ── XLS (BIFF) ─────────────────────────────────────────────────────────────


def extract_xls(ctx: Ctx) -> None:
    import xlrd

    data, _ = ctx.read_bytes()
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        if "encrypt" in str(exc).lower():
            raise EncryptedDocument(str(exc)) from exc
        raise
    try:
        if book.user_name:
            ctx.result.doc_meta["last_modified_by"] = book.user_name
        ctx.result.doc_meta["sheets"] = book.nsheets
        max_rows = ctx.limit("max_rows_per_sheet")
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            title = str(sheet.name)
            ctx.add(title, anchor=ANCHOR_HARD, level=1, role="heading", locator={"sheet": title})

            def rows(sheet: Any = sheet) -> Iterator[tuple[int, list[str]]]:
                for r in range(sheet.nrows):
                    yield r + 1, [_xls_cell(book, cell) for cell in sheet.row(r)]

            _emit_sheet_rows(ctx, _non_empty(rows()), sheet=title, max_rows=max_rows)
            book.unload_sheet(index)
    finally:
        book.release_resources()


def _xls_cell(book: Any, cell: Any) -> str:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return cell_text(xlrd.xldate_as_datetime(cell.value, book.datemode))
        except (ValueError, OverflowError, xlrd.xldate.XLDateError):
            return cell_text(cell.value)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return cell_text(bool(cell.value))
    if cell.ctype in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    return cell_text(cell.value)

