"""PPTX: one hard-anchored block per slide, cited by slide number.

A slide is the unit a person remembers and a citation opens, so its title,
every text frame (inside groups and tables too) and its speaker notes stay in
one block, and no chunk ever spans two slides.
"""

from __future__ import annotations

from typing import Any, Iterator

from app.userdocs.types import ANCHOR_HARD

from ._base import Ctx
from .office import core_properties


def _shape_texts(shapes: Any, skip_id: int | None) -> Iterator[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        # python-pptx hands out a fresh proxy per access, so the title shape
        # is recognised by id, not identity.
        if skip_id is not None and shape.shape_id == skip_id:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes, skip_id)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if any(cells):
                    yield " | ".join(cells)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                yield text


def extract_pptx(ctx: Ctx) -> None:
    from pptx import Presentation

    ctx.result.doc_meta.update(core_properties(ctx))
    source = ctx.source()
    prs = Presentation(source)
    slides = list(prs.slides)
    ctx.result.doc_meta["slides"] = len(slides)
    max_slides = ctx.limit("max_pages")
    for number, slide in enumerate(slides, start=1):
        if number > max_slides:
            ctx.mark_partial()
            break
        title_shape = slide.shapes.title
        parts: list[str] = []
        if title_shape is not None and title_shape.has_text_frame:
            title = title_shape.text_frame.text.strip()
            if title:
                parts.append(title)
        parts.extend(_shape_texts(slide.shapes, title_shape.shape_id if title_shape is not None else None))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                parts.append("Notes: " + notes)
        if parts:
            ctx.add("\n".join(parts), anchor=ANCHOR_HARD, role="para", locator={"slide": number})
